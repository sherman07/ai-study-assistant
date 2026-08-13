import asyncio
import base64
import importlib
import json
import re
import tempfile
import time
import unittest
from types import SimpleNamespace
from pathlib import Path
from urllib.parse import urlparse
from unittest.mock import patch

from fastapi.testclient import TestClient

from backend.app import app
from backend import app as backend_app_module
from backend.core import config as core_config
from backend.app import analyze_materials
from backend.app import build_analysis_fingerprint
from backend.app import build_visual_gallery
from backend.app import choose_learning_depth
from backend.app import enforce_requested_language
from backend.app import file_to_source_unit
from backend.app import finalize_generated_summary
from backend.app import generate_reference_style_multisource_notes
from backend.app import iter_visual_candidates
from backend.app import link_to_source_unit
from backend.app import rebuild_cached_visual_argument_cards
from backend.app import render_pptx_source_preview_svg_images
from backend.app import source_preview_image_url
from backend.app import validate_source_strict_summary
from backend.core.analysis_cache import AnalysisCacheStore
from backend.core import visual_assets
from backend.core.visual_assets import prune_unavailable_visual_markers, visual_asset_url_for_browser

try:
    from pptx import Presentation
except Exception:
    Presentation = None


def assert_served_visual_asset(test_case, url: str, expected_content_type: str = "image/png"):
    parsed = urlparse(url)
    test_case.assertEqual(parsed.scheme, "http")
    test_case.assertEqual(parsed.netloc, "127.0.0.1:8001")
    test_case.assertTrue(parsed.path.startswith("/assets/visuals/"))

    response = TestClient(app).get(parsed.path)
    test_case.assertEqual(response.status_code, 200)
    test_case.assertEqual(response.headers.get("content-type", "").split(";")[0], expected_content_type)


class VisualGalleryTests(unittest.TestCase):
    def test_pptx_svg_visual_rendering_respects_requested_slide_limit(self):
        slide_texts = [
            "Slide 1 table data comparison",
            "Slide 2 chart results",
            "Slide 3 agenda",
            "Slide 4 references",
        ]
        rendered = {
            1: "data:image/png;base64,AA==",
            2: "data:image/png;base64,AA==",
            3: "data:image/png;base64,AA==",
            4: "data:image/png;base64,AA==",
        }

        with (
            patch.object(backend_app_module, "ENABLE_PPTX_SLIDE_RENDER", False, create=True),
            patch.object(backend_app_module, "ENABLE_PPTX_SVG_FALLBACK_RENDER", True, create=True),
            patch.object(backend_app_module, "CONTROLLED_MAX_PPTX_SLIDES_PER_SOURCE", 6, create=True),
            patch.object(backend_app_module, "MAX_VISUAL_IMAGES_PER_SOURCE", 2, create=True),
            patch.object(backend_app_module, "Presentation", return_value=SimpleNamespace(slides=[]), create=True),
            patch.object(backend_app_module, "render_pptx_source_preview_svg_images", return_value=rendered),
            patch.object(backend_app_module, "selected_indices_by_score", return_value=[0, 1, 2, 3]),
        ):
            parts = backend_app_module.render_pptx_slide_screenshots(
                b"pptx bytes",
                "lecture.pptx",
                slide_texts,
                max_slides=2,
            )

        self.assertEqual(len(parts), 4)
        self.assertIn("PPT slide 1", parts[0]["text"])
        self.assertIn("PPT slide 2", parts[2]["text"])
        self.assertNotIn("PPT slide 3", json.dumps(parts))

    def test_pdf_visual_rendering_can_be_disabled_for_a_constrained_host(self):
        with (
            patch.object(backend_app_module, "ENABLE_PDF_VISUAL_EXTRACTION", False, create=True),
            patch("backend.app.extract_pdf", return_value="Lecture text"),
            patch("backend.app.render_pdf_visual_parts", side_effect=AssertionError("PDF rendering should be skipped")),
        ):
            _, source = file_to_source_unit("lecture.pdf", "application/pdf", b"pdf bytes")

        self.assertEqual(source["text_excerpt"], "Lecture text")
        self.assertEqual(source["visual_parts"], [])

    def test_uploaded_image_source_becomes_inline_visual_candidate_without_model_call(self):
        png_data = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )
        _, source = file_to_source_unit("results-chart.png", "image/png", png_data)
        source_units = [source]

        self.assertEqual(source["visual_parts"][0]["type"], "text")
        candidates = iter_visual_candidates(source_units)
        self.assertEqual(len(candidates), 1)
        self.assertFalse(candidates[0]["is_likely_decorative"])
        self.assertIn(candidates[0]["visual_kind"], {"data/table", "graph/chart", "diagram/model", "method/result figure"})

        with patch("backend.app.generate_chat", side_effect=AssertionError("uploaded image fallback should not call the model")):
            cards = rebuild_cached_visual_argument_cards(source_units, "english")
            summary = finalize_generated_summary(
                "## Notes\n\nThe uploaded results chart should be used as source evidence.",
                requested_language="english",
                generation_language="english",
                source_context="",
                source_units=source_units,
                attach_visuals=True,
            )
            gallery = build_visual_gallery(source_units)

        self.assertEqual(len(cards), 1)
        self.assertIn("[[VISUAL:0]]", summary)
        self.assertEqual(len(gallery), 1)
        assert_served_visual_asset(self, gallery[0]["url"], "image/png")

    def test_selected_visual_card_uses_marker_index(self):
        source_units = [{
            "visual_argument_cards": [{
                "index": 17,
                "url": "data:image/png;base64,AA==",
                "title": "Result comparison",
                "caption": "Table with data, results, comparison groups, and statistical evidence.",
                "what_shows": "The table compares results and shows a data pattern.",
                "argument_supported": "The table supports the analysis by showing the comparison directly.",
                "how_to_read": "Read rows, columns, values, and group differences.",
                "visual_kind": "unknown",
            }]
        }]

        gallery = build_visual_gallery(source_units)

        self.assertEqual(len(gallery), 1)
        self.assertEqual(gallery[0]["index"], 0)
        assert_served_visual_asset(self, gallery[0]["url"], "image/png")

    def test_cached_visual_rebuild_does_not_call_model(self):
        source_units = [{
            "display_name": "lecture.pdf",
            "title_candidate": "Lecture data",
            "visual_parts": [
                {
                    "type": "text",
                    "text": (
                        "IN-TEXT SOURCE FIGURE FROM lecture.pdf — PDF page 2. "
                        "Page text preview: table graph data results comparison. "
                        "Image-count=1; drawing-count=12; visual-score=28."
                    ),
                },
                {"type": "image_url", "image_url": {"url": "data:image/png;base64,AA=="}},
            ],
        }]

        with patch("backend.app.generate_chat", side_effect=AssertionError("model should not be called")):
            cards = rebuild_cached_visual_argument_cards(source_units, "english")
            summary = finalize_generated_summary(
                "## Notes\n\nThe table compares the source results.",
                requested_language="english",
                generation_language="english",
                source_context="",
                source_units=source_units,
                attach_visuals=True,
            )
            gallery = build_visual_gallery(source_units)

        self.assertEqual(len(cards), 1)
        self.assertIn("[[VISUAL:0]]", summary)
        self.assertNotIn("This source figure belongs in the notes", summary)
        self.assertNotIn("Figure focus:", summary)
        self.assertEqual(len(gallery), 1)
        assert_served_visual_asset(self, gallery[0]["url"], "image/png")

    def test_finalize_keeps_visual_cards_out_of_main_summary_and_is_idempotent(self):
        source_units = [{
            "visual_argument_cards": [{
                "index": 0,
                "url": "data:image/png;base64,AA==",
                "title": "Result table",
                "caption": "Table with data and results.",
                "what_shows": "The table compares results and shows a data pattern.",
                "argument_supported": "The table supports the analysis by showing the comparison directly.",
                "how_to_read": "Read rows, columns, values, and group differences.",
                "visual_kind": "data/table",
            }]
        }]
        polluted = (
            "# Generated Study Notes\n\n"
            "## Overview\n\n"
            "This lecture teaches how to interpret the result table as evidence.\n\n"
            "## Source Examples and Evidence\n\n"
            "### Result table\n\n"
            "This source figure belongs in the notes because it shows: Result table. "
            "It supports the surrounding explanation by making the method, pattern, mechanism, or contrast visible "
            "rather than asking the student to memorise an abstract claim.\n\n"
            "*Figure focus: Result table. It supports the surrounding explanation by making the method, pattern, mechanism, or contrast visible.*\n\n"
            "[[VISUAL:0]]\n\n"
            "## Core Argument\n\n"
            "The real study task is to explain what the evidence can and cannot prove."
        )

        once = finalize_generated_summary(
            polluted,
            requested_language="english",
            generation_language="english",
            source_context="",
            source_units=source_units,
            attach_visuals=True,
        )
        twice = finalize_generated_summary(
            once,
            requested_language="english",
            generation_language="english",
            source_context="",
            source_units=source_units,
            attach_visuals=True,
        )
        gallery = build_visual_gallery(source_units)

        self.assertEqual(once, twice)
        for bad_phrase in (
            "Source Examples and Evidence",
            "This source figure belongs in the notes",
            "method, pattern, mechanism, or contrast visible",
            "Figure focus:",
        ):
            self.assertNotIn(bad_phrase, twice)
        self.assertEqual(twice.count("[[VISUAL:0]]"), 1)
        self.assertIn("## Overview", twice)
        self.assertIn("## Core Argument", twice)
        self.assertIn("The real study task is to explain", twice)
        self.assertEqual(len(gallery), 1)

    def test_visual_gallery_serves_runtime_asset_url(self):
        source_units = [{
            "visual_argument_cards": [{
                "index": 0,
                "url": "data:image/png;base64,AA==",
                "title": "Source data table",
                "caption": "Table with data, results, comparison groups, and statistical evidence.",
                "what_shows": "The table compares results and shows a data pattern.",
                "argument_supported": "The table supports the analysis by showing the comparison directly.",
                "how_to_read": "Read rows, columns, values, and group differences.",
                "visual_kind": "data/table",
            }]
        }]

        gallery = build_visual_gallery(source_units)

        self.assertEqual(len(gallery), 1)
        self.assertFalse(gallery[0]["url"].startswith("data:"))
        assert_served_visual_asset(self, gallery[0]["url"], "image/png")

    def test_visual_asset_mirrors_new_data_url_to_durable_storage(self):
        png_data = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )
        data_url = "data:image/png;base64," + base64.b64encode(png_data).decode("ascii")

        with (
            patch.object(visual_assets, "durable_visual_storage_enabled", return_value=True),
            patch.object(visual_assets, "persist_visual_asset_to_durable_storage", return_value=True) as persist,
        ):
            url = visual_asset_url_for_browser(data_url)

        self.assertTrue(url.endswith(".png"))
        persist.assert_called_once()
        asset_name, content_type, stored_data = persist.call_args.args
        self.assertTrue(asset_name.endswith(".png"))
        self.assertEqual(content_type, "image/png")
        self.assertEqual(stored_data, png_data)

    def test_visual_asset_route_recovers_missing_runtime_file_from_durable_storage(self):
        asset_name = "durable-test-visual.png"
        asset_path = backend_app_module.RUNTIME_ASSETS_DIR / "visuals" / asset_name
        asset_path.unlink(missing_ok=True)
        png_data = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )

        try:
            with patch.object(
                backend_app_module,
                "fetch_visual_asset_from_durable_storage",
                return_value=(png_data, "image/png"),
            ) as fetch:
                response = TestClient(app).get(f"/assets/visuals/{asset_name}")
        finally:
            asset_path.unlink(missing_ok=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers.get("content-type"), "image/png")
        self.assertEqual(response.content, png_data)
        fetch.assert_called_once_with(asset_name)

    def test_source_preview_image_url_serves_runtime_asset_when_requested(self):
        png_data = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )

        url = source_preview_image_url(png_data, "image/png", browser_asset=True)

        self.assertFalse(url.startswith("data:"))
        assert_served_visual_asset(self, url, "image/png")

    def test_visual_asset_url_preserves_existing_browser_url(self):
        url = "http://127.0.0.1:8001/assets/visuals/already.png"

        self.assertEqual(visual_asset_url_for_browser(url), url)

    def test_visual_asset_url_rejects_invalid_or_non_raster_data_urls(self):
        invalid_url = "data:image/png;base64,not-valid-base64"
        svg_url = "data:image/svg+xml;base64,PHN2Zy8+"

        self.assertEqual(visual_asset_url_for_browser(invalid_url), "")
        self.assertEqual(visual_asset_url_for_browser(svg_url), "")

    def test_visual_asset_url_rejects_local_filesystem_paths(self):
        self.assertEqual(visual_asset_url_for_browser("file:///Users/student/source/page.png"), "")
        self.assertEqual(visual_asset_url_for_browser("/Users/student/source/page.png"), "")
        self.assertEqual(visual_asset_url_for_browser("C:\\Users\\student\\source\\page.png"), "")

    def test_visual_gallery_rejects_local_filesystem_card_urls(self):
        source_units = [{
            "visual_argument_cards": [{
                "index": 0,
                "url": "/Users/student/source/page.png",
                "title": "Local-only table",
                "caption": "Table with data, results, comparison groups, and statistical evidence.",
                "what_shows": "The table compares results and shows a data pattern.",
                "argument_supported": "The table supports the analysis by showing the comparison directly.",
                "how_to_read": "Read rows, columns, values, and group differences.",
                "visual_kind": "data/table",
            }]
        }]

        gallery = build_visual_gallery(source_units)

        self.assertEqual(gallery, [])

    def test_cached_id_only_visual_metadata_preserves_matching_marker(self):
        summary = "## Notes\n\nUse the evidence table here.\n\n[[VISUAL:7]]\n\nNext point."
        items = [{
            "id": 7,
            "url": "http://127.0.0.1:8001/assets/visuals/id-only.png",
            "title": "ID-only source figure",
        }]

        pruned = prune_unavailable_visual_markers(summary, items)

        self.assertIn("[[VISUAL:7]]", pruned)

    def test_visual_markers_are_reindexed_after_unrenderable_cards_are_removed(self):
        source_units = [{
            "visual_argument_cards": [
                {
                    "index": 0,
                    "url": "",
                    "title": "Missing visual",
                    "caption": "Table with data, results, comparison groups, and statistical evidence.",
                    "what_shows": "The unavailable table compares results.",
                    "argument_supported": "This card should not get a marker because it has no image URL.",
                    "how_to_read": "Read rows, columns, values, and group differences.",
                    "visual_kind": "data/table",
                },
                {
                    "index": 1,
                    "url": "data:image/png;base64,AA==",
                    "title": "Renderable result table",
                    "caption": "Table with data, results, comparison groups, and statistical evidence.",
                    "what_shows": "The table compares results and shows a data pattern.",
                    "argument_supported": "The table supports the analysis by showing the comparison directly.",
                    "how_to_read": "Read rows, columns, values, and group differences.",
                    "visual_kind": "data/table",
                },
            ]
        }]

        summary = finalize_generated_summary(
            "## Notes\n\nThe table compares source results and supports the key claim.",
            requested_language="english",
            generation_language="english",
            source_context="",
            source_units=source_units,
            attach_visuals=True,
        )
        gallery = build_visual_gallery(source_units)

        self.assertIn("[[VISUAL:0]]", summary)
        self.assertNotIn("[[VISUAL:1]]", summary)
        self.assertEqual(len(gallery), 1)
        self.assertEqual(gallery[0]["index"], 0)
        self.assertEqual(gallery[0]["title"], "Renderable result table")
        assert_served_visual_asset(self, gallery[0]["url"], "image/png")

    def test_pptx_svg_fallback_is_rasterized_for_runtime_assets(self):
        if Presentation is None:
            self.skipTest("python-pptx is not installed")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(914400, 914400, 5486400, 914400)
        textbox.text = "Results table and comparison graph"

        rendered = render_pptx_source_preview_svg_images(prs, 1)

        self.assertIn(1, rendered)
        self.assertTrue(rendered[1].startswith("data:image/png;base64,"))
        assert_served_visual_asset(self, visual_asset_url_for_browser(rendered[1]), "image/png")

    def test_pptx_svg_browser_preview_uses_runtime_asset_url(self):
        if Presentation is None:
            self.skipTest("python-pptx is not installed")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(914400, 914400, 5486400, 914400)
        textbox.text = "Results table and comparison graph"

        rendered = render_pptx_source_preview_svg_images(prs, 1, browser_assets=True)

        self.assertIn(1, rendered)
        self.assertFalse(rendered[1].startswith("data:"))
        assert_served_visual_asset(self, rendered[1], "image/png")

    def test_youtube_link_preserves_frame_visual_parts_for_inline_candidates(self):
        frame_parts = [
            {
                "type": "text",
                "text": (
                    "IN-TEXT SOURCE FIGURE FROM lecture video — video frame 1 sampled at approximately 00:10. "
                    "The frame shows a graph with data, results, comparison groups, and visual-score=32."
                ),
            },
            {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64,AA=="}},
        ]
        meta = {
            "url": "https://www.youtube.com/watch?v=dQw4w9WgXcQ",
            "source_identity": "youtube:dQw4w9WgXcQ",
            "detected_title": "Lecture video",
            "content_hash": "hash",
            "visual_parts": frame_parts,
        }

        with patch("backend.app.analyse_youtube_url", return_value=("Transcript text", frame_parts, meta)):
            parts, source = link_to_source_unit("https://www.youtube.com/watch?v=dQw4w9WgXcQ")

        self.assertEqual(source["visual_parts"], frame_parts)
        self.assertEqual(parts[-1]["image_url"]["url"], "data:image/jpeg;base64,AA==")
        candidates = iter_visual_candidates([source])
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["url"], "data:image/jpeg;base64,AA==")


class SourceIdentityTests(unittest.TestCase):
    def test_uploaded_file_identity_keeps_raw_bytes_when_extracted_text_matches(self):
        with patch("backend.app.extract_pdf", return_value=""), patch("backend.app.render_pdf_visual_parts", return_value=[]):
            _, first = file_to_source_unit("scan-a.pdf", "application/pdf", b"pdf bytes a")
            _, second = file_to_source_unit("scan-b.pdf", "application/pdf", b"pdf bytes b")

        self.assertNotEqual(first["source_identity"], second["source_identity"])
        self.assertNotEqual(first["content_hash"], second["content_hash"])
        self.assertNotEqual(
            build_analysis_fingerprint("auto", [first], "detailed", "professor_mode"),
            build_analysis_fingerprint("auto", [second], "detailed", "professor_mode"),
        )




if __name__ == "__main__":
    unittest.main()
